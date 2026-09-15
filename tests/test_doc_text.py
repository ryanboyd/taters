"""
Document reading: the folder source's new reach (.txt, .docx, .doc, .pdf).

The resilience contract matters more than the happy paths: a broken file --
the right extension wrapped around an image, a truncated PDF, a legacy .doc
with no converter -- costs only itself. It is treated as a document with no
text: skipped with a warning naming it, never allowed to break a run.
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pytest

from taters.helpers.doc_text import (
    DOCUMENT_PATTERN,
    DOCUMENT_SUFFIXES,
    DocumentReadError,
    read_document_text,
)
from taters.helpers.text_gather import txt_folder_to_analysis_ready_csv

#: A complete, valid one-page PDF whose text layer says "Hello pdf world".
#: pypdf reconstructs the (deliberately lazy) xref, so this stays readable.
MINIMAL_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 44>>stream
BT /F1 12 Tf 72 720 Td (Hello pdf world) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
startxref
0
%%EOF"""

#: Bytes that are unmistakably not text (a PNG header carries NULs).
BINARY = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + b"\x00" * 64


def _docx(path, paragraphs=("Hello docx world.",), table=None):
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=1, cols=len(table))
        for i, cell in enumerate(table):
            t.rows[0].cells[i].text = cell
    d.save(str(path))
    return path


def _blank_pdf(path):
    import pypdf

    w = pypdf.PdfWriter()
    w.add_blank_page(width=612, height=792)
    with open(path, "wb") as f:
        w.write(f)
    return path


# ---------------------------------------------------------------------------
# The readers
# ---------------------------------------------------------------------------

def test_docx_text_includes_paragraphs_and_tables(tmp_path):
    """Table text lives outside `paragraphs`; a Word doc that is mostly a
    table would otherwise come back empty."""
    f = _docx(tmp_path / "doc.docx", table=("cell one", "cell two"))
    text = read_document_text(f)
    assert "Hello docx world." in text
    assert "cell one" in text and "cell two" in text


def test_pdf_text_layer_is_extracted(tmp_path):
    f = tmp_path / "doc.pdf"
    f.write_bytes(MINIMAL_PDF)
    assert "Hello pdf world" in read_document_text(f)


def test_a_pdf_without_a_text_layer_has_zero_text(tmp_path):
    """No OCR, by design: a scanned page is a document with no text."""
    assert read_document_text(_blank_pdf(tmp_path / "scan.pdf")) == ""


def test_a_custom_text_suffix_still_reads_as_plain_text(tmp_path):
    """A user's pattern (*.md, *.log) is their statement that the files are
    text; refusing unknown suffixes broke exactly that."""
    f = tmp_path / "notes.md"
    f.write_text("markdown words", encoding="utf-8")
    assert read_document_text(f) == "markdown words"


@pytest.mark.parametrize("name", ["fake.docx", "fake.pdf", "fake.txt"])
def test_a_binary_file_in_document_clothing_is_refused_not_garbled(tmp_path, name):
    """
    The requested contract: the right extension around an image must not
    break a run OR contribute byte soup. Every reader turns it into a named
    refusal the gatherer converts to 'no text'.
    """
    f = tmp_path / name
    f.write_bytes(BINARY)
    with pytest.raises(DocumentReadError, match=name):
        read_document_text(f)


def test_a_legacy_doc_without_antiword_says_how_to_fix_it(tmp_path, monkeypatch):
    import taters.helpers.doc_text as dt

    monkeypatch.setattr(dt.shutil, "which", lambda _name: None)
    f = tmp_path / "old.doc"
    f.write_bytes(BINARY)
    with pytest.raises(DocumentReadError, match="Save As"):
        read_document_text(f)


# ---------------------------------------------------------------------------
# The gatherer: a broken file costs only itself
# ---------------------------------------------------------------------------

def _mixed_folder(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.txt").write_text("plain words here", encoding="utf-8")
    _docx(root / "b.docx", paragraphs=("word words",))
    (root / "c.pdf").write_bytes(MINIMAL_PDF)
    return root


def test_a_mixed_folder_gathers_every_document_type(tmp_path):
    root = _mixed_folder(tmp_path)
    out = txt_folder_to_analysis_ready_csv(
        root_dir=root, out_csv=tmp_path / "out.csv", recursive=True)

    import csv
    rows = {r["text_id"]: r["text"] for r in
            csv.DictReader(open(out, encoding="utf-8-sig"))}
    assert set(rows) == {"a", "b", "c"}
    assert rows["a"] == "plain words here"
    assert "word words" in rows["b"]
    assert "Hello pdf world" in rows["c"]


def test_broken_documents_cost_only_themselves(tmp_path):
    """An image wearing a .docx name, a scan with no text layer: each is a
    warning naming the file, and every readable neighbor still lands."""
    root = _mixed_folder(tmp_path)
    (root / "imposter.docx").write_bytes(BINARY)
    _blank_pdf(root / "scan.pdf")

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        out = txt_folder_to_analysis_ready_csv(
            root_dir=root, out_csv=tmp_path / "out.csv", recursive=True)

    import csv
    rows = [r["text_id"] for r in csv.DictReader(open(out, encoding="utf-8-sig"))]
    assert rows == ["a", "b", "c"]
    messages = [str(w.message) for w in caught]
    assert any("imposter.docx" in m for m in messages)
    assert any("scan.pdf" in m and "no machine-readable text" in m
               for m in messages)


def test_multiple_patterns_walk_in_sorted_order(tmp_path):
    """With several globs the concatenation order would otherwise depend on
    the filesystem, and two runs of one folder must write the same file."""
    root = tmp_path / "docs"
    root.mkdir()
    for name in ("zeta.txt", "alpha.pdf", "mid.docx"):
        if name.endswith(".txt"):
            (root / name).write_text("z words", encoding="utf-8")
        elif name.endswith(".pdf"):
            (root / name).write_bytes(MINIMAL_PDF)
        else:
            _docx(root / name)

    import csv
    out = txt_folder_to_analysis_ready_csv(
        root_dir=root, out_csv=tmp_path / "out.csv")
    rows = [r["text_id"] for r in csv.DictReader(open(out, encoding="utf-8-sig"))]
    assert rows == ["alpha", "mid", "zeta"]


def test_the_document_constants_agree_with_each_other():
    globbed = {g.strip().lstrip("*") for g in DOCUMENT_PATTERN.split(";")}
    assert globbed == set(DOCUMENT_SUFFIXES)


def test_gathering_is_its_own_phase_with_named_inflight_documents(tmp_path):
    """
    Two reports shaped this: a frozen "reading the input" bar (now: per-file
    progress), and a step label that looked like it was doing someone else's
    work (now: the phase says "gathering documents" and each document being
    read gets its own entry, the way ffmpeg's per-file rows do).
    """
    root = _mixed_folder(tmp_path)
    calls = []
    txt_folder_to_analysis_ready_csv(
        root_dir=root, out_csv=tmp_path / "out.csv", recursive=True, workers=1,
        on_progress=lambda done, total, message=None, unit=None, **extra:
            calls.append((done, total, message, extra.get("inflight"))))

    assert all(c[2] == "gathering documents" for c in calls)
    consumed = [(c[0], c[1]) for c in calls if c[3] == []]
    assert consumed[-1] == (3, 3)
    seen_inflight = {name for c in calls for name in (c[3] or [])}
    assert seen_inflight == {"a.txt", "b.docx", "c.pdf"}


def test_a_four_argument_sink_still_works_without_inflight(tmp_path):
    """The documented on_progress contract is four arguments; the in-flight
    detail goes only to sinks that can accept it."""
    root = _mixed_folder(tmp_path)
    calls = []
    txt_folder_to_analysis_ready_csv(
        root_dir=root, out_csv=tmp_path / "out.csv", recursive=True,
        on_progress=lambda done, total, message=None, unit=None:
            calls.append((done, total, message)))
    assert (3, 3, "gathering documents") in calls


def test_every_text_analyzer_threads_progress_into_the_gather(tmp_path):
    """The wiring, not just the mechanism: analyze_readability's txt path must
    hand its own on_progress down, or the fix helps nobody."""
    from taters.text.analyze_readability import analyze_readability

    root = _mixed_folder(tmp_path)
    calls = []
    analyze_readability(
        txt_dir=root, gathered_csv=tmp_path / "g.csv",
        out_features_csv=tmp_path / "f.csv", overwrite_existing=True,
        on_progress=lambda done, total, message=None, unit=None:
            calls.append((done, total, message)))

    assert any(m == "gathering documents" and t == 3
               for _d, t, m in calls)


def test_a_grumbly_pdf_becomes_one_warning_naming_the_file(tmp_path, caplog):
    """
    pypdf logs dozens of "Ignoring wrong pointing object..." lines straight to
    stderr for a slightly malformed PDF -- a wall of noise naming no file.
    Suppressed and counted instead: one warning, with the filename, and the
    text still extracted.
    """
    import logging

    f = tmp_path / "grumbly.pdf"
    f.write_bytes(MINIMAL_PDF)      # the lazy xref makes pypdf grumble

    with caplog.at_level(logging.WARNING, logger="pypdf"):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            text = read_document_text(f)

    assert "Hello pdf world" in text
    assert not caplog.records, "pypdf's log lines must not reach the screen"
    messages = [str(w.message) for w in caught]
    assert any("grumbly.pdf" in m and "malformed" in m for m in messages)


def test_surrogate_halves_from_a_pdf_cannot_poison_the_csv_write(monkeypatch):
    """
    From a real run: pypdf extracted a math glyph as UTF-16 surrogates, and
    the gathered-CSV write died minutes later with "surrogates not allowed".
    Paired surrogates are rescued into the character they encode; unpaired
    ones become U+FFFD; either way the text encodes.
    """
    import taters.helpers.doc_text as dt

    # a proper pair (𝑥 the way pypdf sometimes hands it over: two surrogate
    # code points, not the actual character) and then a lone half.
    monkeypatch.setattr(dt, "_pdf_text",
                        lambda path: "the \ud835\udc65 variable")
    paired = read_document_text(Path("whatever.pdf"))
    assert paired == "the \U0001d465 variable"
    paired.encode("utf-8")

    monkeypatch.setattr(dt, "_pdf_text", lambda path: "broken \ud835 half")
    lone = read_document_text(Path("whatever.pdf"))
    assert "\ud835" not in lone
    lone.encode("utf-8")


def test_clean_text_passes_the_scrub_untouched():
    from taters.helpers.doc_text import _scrub_surrogates

    text = "naïve café 猫 :-) straße 𝑥"
    assert _scrub_surrogates(text) is text


def test_an_enormous_document_is_flagged_at_gather_time(tmp_path):
    """A proceedings volume as one 'document' is legal but makes downstream
    scoring crawl; naming it now beats looking hung twenty minutes later."""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "normal.txt").write_text("ordinary words " * 100, encoding="utf-8")
    (root / "tome.txt").write_text("word " * 120_000, encoding="utf-8")

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        txt_folder_to_analysis_ready_csv(
            root_dir=root, out_csv=tmp_path / "out.csv")

    messages = [str(w.message) for w in caught]
    assert any("tome.txt" in m and "very large" in m for m in messages)
    assert not any("normal.txt" in m for m in messages)


# ---------------------------------------------------------------------------
# Parallel reading: faster, byte-identical
# ---------------------------------------------------------------------------

def test_parallel_gather_output_is_byte_identical_to_serial(tmp_path):
    """
    The determinism contract: one worker or many, the gathered CSV is the
    same file. No binning needed -- files are walked sorted and results are
    consumed in submission order.
    """
    root = _mixed_folder(tmp_path)
    (root / "zz_extra.txt").write_text("tail words", encoding="utf-8")

    serial = tmp_path / "serial.csv"
    parallel = tmp_path / "parallel.csv"
    txt_folder_to_analysis_ready_csv(root_dir=root, out_csv=serial,
                                     recursive=True, workers=1)
    txt_folder_to_analysis_ready_csv(root_dir=root, out_csv=parallel,
                                     recursive=True, workers=3)

    assert serial.read_bytes() == parallel.read_bytes()


def test_worker_warnings_survive_the_process_boundary(tmp_path):
    """Skip-and-warn must be identical under any worker count: a broken file
    read inside a worker still warns, by name, in the parent."""
    root = _mixed_folder(tmp_path)
    (root / "imposter.docx").write_bytes(BINARY)

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        out = txt_folder_to_analysis_ready_csv(
            root_dir=root, out_csv=tmp_path / "out.csv",
            recursive=True, workers=2)

    import csv
    rows = [r["text_id"] for r in csv.DictReader(open(out, encoding="utf-8-sig"))]
    assert rows == ["a", "b", "c"]
    assert any("imposter.docx" in str(w.message) for w in caught)


def test_resolve_workers_contract(monkeypatch):
    """0 = automatic (three-quarters of the cores, "leave some for the
    human"); explicit asks are honored only up to the machine's logical
    core count -- past the cores, more processes is strictly worse."""
    from taters.helpers import parallel_map as pm

    monkeypatch.setattr(pm.os, "cpu_count", lambda: 8)
    assert pm.max_workers() == 8
    assert pm.resolve_workers(1) == 1
    assert pm.resolve_workers(3) == 3
    assert pm.resolve_workers(-5) == 1
    assert pm.resolve_workers(0) == 6, "auto = 8 - 8//4"
    assert pm.resolve_workers(8) == 8
    assert pm.resolve_workers(32) == 8, "an ask past the cores is clamped"

    # small machines: the quarter rounds down before it can starve anyone.
    monkeypatch.setattr(pm.os, "cpu_count", lambda: 2)
    assert pm.resolve_workers(0) == 2
    monkeypatch.setattr(pm.os, "cpu_count", lambda: 4)
    assert pm.resolve_workers(0) == 3
    monkeypatch.setattr(pm.os, "cpu_count", lambda: 1)
    assert pm.resolve_workers(0) == 1
    monkeypatch.setattr(pm.os, "cpu_count", lambda: None)
    assert pm.resolve_workers(0) == 1
    assert pm.resolve_workers(5) == 1, "cpu_count unknown: the ceiling is 1"


def test_in_worker_warnings_ride_home_with_their_results(tmp_path):
    """The malformed-PDF notice is raised *inside* a worker process; it must
    be re-issued in the parent, or parallel runs silently lose warnings that
    serial runs show. (Broken-file errors travel a different path.)"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "grumbly.pdf").write_bytes(MINIMAL_PDF)   # lazy xref -> grumbles

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        txt_folder_to_analysis_ready_csv(
            root_dir=root, out_csv=tmp_path / "out.csv", workers=2)

    assert any("grumbly.pdf" in str(w.message) and "malformed" in str(w.message)
               for w in caught)


def test_work_actually_crosses_the_process_boundary():
    """workers=2 must genuinely use other processes -- an implementation that
    quietly fell back to inline would pass every determinism test."""
    import os

    from taters.helpers import parallel_map

    pids = set(parallel_map.ordered_parallel_map(
        lambda x: os.getpid(), range(6), workers=2,
        pool_fn=parallel_map._worker_pid))
    assert os.getpid() not in pids


def test_an_unspawnable_main_degrades_to_serial_instead_of_crashing(monkeypatch):
    """A REPL or stdin-script parent cannot be re-imported by spawn: every
    worker would die on startup (found the hard way). Those callers get the
    serial path -- slow, never broken."""
    import os

    from taters.helpers import parallel_map

    monkeypatch.setattr(parallel_map, "_spawn_safe", lambda: False)
    pids = set(parallel_map.ordered_parallel_map(
        lambda x: os.getpid(), range(4), workers=4,
        pool_fn=parallel_map._worker_pid))
    assert pids == {os.getpid()}


def test_inflight_names_flow_from_real_worker_pools(tmp_path):
    """The lifecycle hooks must fire on the *pool* path too -- submission and
    completion happen in the parent and an executor thread respectively, and
    dropping either would quietly blank the sub-bars only when parallel."""
    root = _mixed_folder(tmp_path)
    calls = []
    txt_folder_to_analysis_ready_csv(
        root_dir=root, out_csv=tmp_path / "out.csv", recursive=True, workers=2,
        on_progress=lambda done, total, message=None, unit=None, **extra:
            calls.append(list(extra.get("inflight") or [])))

    seen = {name for names in calls for name in names}
    assert {"a.txt", "b.docx", "c.pdf"} <= seen
    assert calls[-1] == [], "everything must be retired by the end"
